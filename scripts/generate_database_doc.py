from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


TABLE_META = {
    "roles": (
        "Lưu danh sách vai trò trong hệ thống như chủ nhà hàng, quản lý, thu ngân, phục vụ và bếp.",
        "Khi tạo tài khoản nhân sự mới, hệ thống gán role để quyết định phạm vi chức năng được phép truy cập.",
    ),
    "permissions": (
        "Lưu danh sách quyền thao tác trong hệ thống và hỗ trợ phân cấp quyền cha – con.",
        "Một nhóm quyền như USER có thể chứa các quyền con USER_CREATE, USER_READ, USER_UPDATE, USER_DELETE.",
    ),
    "role_permissions": (
        "Bảng nối nhiều – nhiều giữa vai trò và quyền.",
        "Vai trò MANAGER có thể được gán nhiều quyền; mỗi quyền cũng có thể gán cho nhiều vai trò.",
    ),
    "users": (
        "Lưu tài khoản đăng nhập của nhân sự sử dụng hệ thống.",
        "Thu ngân, quản lý hoặc phục vụ dùng tài khoản riêng để đăng nhập API và thực hiện nghiệp vụ.",
    ),
    "table_areas": (
        "Chia khu vực bàn như tầng 1, sân vườn, phòng VIP.",
        "Nhân viên cần xem bàn theo khu vực thay vì một danh sách quá dài.",
    ),
    "restaurant_tables": (
        "Lưu từng bàn ăn thực tế trong nhà hàng.",
        "Nhân viên chọn bàn khi mở phiên phục vụ hoặc nhận order.",
    ),
    "table_sessions": (
        "Đại diện cho một lượt khách ngồi tại một bàn.",
        "Cùng một bàn trong ngày có thể phục vụ nhiều lượt khách khác nhau; mỗi lượt phải có session riêng.",
    ),
    "item_types": (
        "Phân loại bản chất sản phẩm như đồ ăn, đồ uống, combo, topping, phụ thu.",
        "Giúp lọc và quản lý dữ liệu theo bản chất nghiệp vụ thay vì chỉ theo category hiển thị.",
    ),
    "cooking_methods": (
        "Lưu phương pháp chế biến như xào, chiên, nướng, luộc.",
        "Giúp lọc món theo cách chế biến và hỗ trợ báo cáo hoặc hiển thị thông minh.",
    ),
    "menu_categories": (
        "Tổ chức danh mục menu theo cấu trúc cha – con.",
        "Đồ uống có thể chia thành trà sữa, nước đóng chai; đồ ăn chia thành món chính, món nướng.",
    ),
    "menu_items": (
        "Lưu món gốc ở mức nghiệp vụ, chưa phải SKU bán ra cuối cùng.",
        "Một món có thể có nhiều biến thể; menu_items giữ phần mô tả chung của món.",
    ),
    "menu_item_variants": (
        "Lưu biến thể bán ra thực tế và mức giá hiện hành của từng biến thể.",
        "Cùng một món trà sữa có thể có size M và size L với giá khác nhau.",
    ),
    "option_groups": (
        "Định nghĩa nhóm tùy chọn như mức đường, mức đá, topping.",
        "Một món có thể cần chọn 1 mức đường nhưng được chọn nhiều topping.",
    ),
    "option_values": (
        "Lưu từng giá trị cụ thể trong mỗi nhóm tùy chọn.",
        "Nhóm mức đá có 0%, 50%, 100%; nhóm topping có pudding, kem cheese.",
    ),
    "variant_option_groups": (
        "Bảng nối biến thể món với các nhóm option được phép áp dụng.",
        "Trà sữa dùng được mức đường, đá, topping; nước lon có thể không có option nào.",
    ),
    "combos": (
        "Lưu thông tin combo bán ra.",
        "Nhà hàng có combo trưa, combo văn phòng hoặc combo theo chương trình khuyến mãi.",
    ),
    "combo_groups": (
        "Chia combo thành các nhóm chọn bên trong.",
        "Một combo có nhóm chọn món chính và nhóm chọn nước uống.",
    ),
    "combo_group_items": (
        "Xác định món hoặc biến thể nào được phép xuất hiện trong từng nhóm combo.",
        "Khách chọn Coca hoặc Trà đào trong nhóm nước; mỗi lựa chọn có thể kèm phụ thu.",
    ),
    "cart_orders": (
        "Lưu order tạm của một bàn trong lúc phục vụ.",
        "Khách đang ngồi tại bàn có thể gọi thêm món nhiều lần trước khi thanh toán.",
    ),
    "cart_order_items": (
        "Lưu từng dòng món trong order tạm.",
        "Một order có nhiều món, mỗi món có biến thể, option, ghi chú và đơn giá snapshot riêng.",
    ),
    "cart_order_item_options": (
        "Lưu các option đã chọn cho từng dòng món trong order tạm.",
        "Một ly trà sữa có thể chọn 50% đường, 50% đá và thêm pudding.",
    ),
    "invoices": (
        "Lưu hóa đơn chính thức khi khách thanh toán.",
        "Khi bàn yêu cầu tính tiền, hệ thống chuyển dữ liệu từ order tạm sang hóa đơn để khóa doanh thu.",
    ),
    "invoice_items": (
        "Lưu snapshot từng dòng món sau khi hóa đơn được tạo.",
        "Giữ nguyên lịch sử doanh thu ngay cả khi menu đổi tên, đổi giá hoặc ngừng bán sau đó.",
    ),
    "payments": (
        "Lưu giao dịch thanh toán của hóa đơn.",
        "Nhà hàng có thể nhận tiền mặt, QR, thẻ hoặc ví điện tử; mỗi lần thu tiền nên có bản ghi riêng.",
    ),
}

TARGET_TABLES = list(TABLE_META.keys())


def normalize_type(line: str, method: str) -> str:
    null_suffix = ", NULL" if "nullable()" in line else ""
    extra = []
    if "unique()" in line:
        extra.append("UNIQUE")
    if "primary([" in line or "->primary(" in line:
        extra.append("PK")
    if "constrained(" in line or method == "foreignId":
        extra.append("FK")

    mapping = {
        "id": "BIGINT, PK",
        "foreignId": "BIGINT",
        "string": "VARCHAR",
        "char": "CHAR",
        "integer": "INT",
        "decimal": "DECIMAL",
        "boolean": "BOOLEAN",
        "text": "TEXT",
        "enum": "ENUM",
        "dateTime": "DATETIME",
        "timestamp": "DATETIME",
        "time": "TIME",
        "rememberToken": "VARCHAR, NULL",
    }
    base = mapping.get(method, method.upper())
    if extra:
        base = f"{base}, " + ", ".join(extra)
    return base + null_suffix


def explain_column(column: str) -> str:
    mapping = {
        "id": "Khóa chính của bản ghi.",
        "is_active": "Cờ bật/tắt để ngừng sử dụng mềm mà không xóa cứng dữ liệu.",
        "created_at": "Ngày tạo bản ghi.",
        "updated_at": "Ngày cập nhật gần nhất.",
        "code": "Mã ổn định để tra cứu và dùng trong code.",
        "name": "Tên hiển thị cho người dùng.",
        "remark": "Ghi chú hoặc mô tả ngắn.",
        "status": "Trạng thái nghiệp vụ của bản ghi.",
        "sort_order": "Thứ tự hiển thị trên giao diện.",
    }
    if column in mapping:
        return mapping[column]
    if column.endswith("_id"):
        return "Khóa ngoại liên kết sang bảng liên quan."
    if column.endswith("_amount"):
        return "Giá trị tiền tệ dùng cho tính toán nghiệp vụ."
    if column.endswith("_price"):
        return "Giá trị giá bán, phụ thu hoặc giá vốn."
    if column.endswith("_at") or column.endswith("_time"):
        return "Thời điểm phát sinh hoặc cập nhật nghiệp vụ."
    if column.endswith("_snapshot"):
        return "Dữ liệu snapshot nhằm giữ nguyên lịch sử tại thời điểm giao dịch."
    return "Thuộc tính nghiệp vụ của bản ghi."


def example_note(column: str) -> str:
    samples = {
        "code": "Ví dụ: AREA_1, TS_TT, CB_TRUA_A",
        "status": "Ví dụ: ACTIVE, OPEN, PAID",
        "is_active": "Y = đang dùng, N = ngừng dùng",
        "sort_order": "Ví dụ: 1, 2, 3",
        "guest_count": "Ví dụ: 2, 4, 6 khách",
        "capacity": "Ví dụ: bàn 4 người, bàn 6 người",
        "price": "Ví dụ: 39.000, 49.000",
        "total_amount": "Ví dụ: 92.000",
        "order_no": "Ví dụ: ORD20260320-0001",
        "no": "Ví dụ: INV-20260320-000001",
    }
    return samples.get(column, "")


def parse_migration(path: Path):
    text = path.read_text(encoding="utf-8")
    match = re.search(r"Schema::create\('([^']+)'", text)
    if not match:
        return None
    table_name = match.group(1)
    if table_name not in TARGET_TABLES:
        return None

    columns = []
    for line in text.splitlines():
        line = line.strip()
        if "$table->timestamps();" in line:
            columns.append(("created_at", "DATETIME", explain_column("created_at"), ""))
            columns.append(("updated_at", "DATETIME", explain_column("updated_at"), ""))
            continue
        if "$table->rememberToken();" in line:
            columns.append(("remember_token", "VARCHAR, NULL", "Token ghi nhớ đăng nhập nếu dùng.", ""))
            continue

        col = re.search(r"\$table->(id|foreignId|string|char|integer|decimal|boolean|text|enum|dateTime|timestamp|time)\('([^']+)'", line)
        if not col:
            continue

        method, column = col.groups()
        columns.append((column, normalize_type(line, method), explain_column(column), example_note(column)))

    if table_name == "role_permissions":
        columns = [
            ("role_id", "BIGINT, PK, FK", "Vai trò được gán quyền.", "Ví dụ: MANAGER"),
            ("permission_id", "BIGINT, PK, FK", "Quyền cụ thể được gán.", "Ví dụ: USER_READ"),
        ]

    return table_name, columns


def build_document() -> Document:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)
    styles["Heading 1"].font.name = "Times New Roman"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.name = "Times New Roman"
    styles["Heading 2"].font.size = Pt(13)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("TÀI LIỆU MÔ TẢ CƠ SỞ DỮ LIỆU HỆ THỐNG QUẢN LÝ NHÀ HÀNG")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Tài liệu tổng hợp từ schema hiện tại của dự án").italic = True

    doc.add_heading("1. Phạm vi tài liệu", level=1)
    doc.add_paragraph(
        "Tài liệu này mô tả các bảng dữ liệu nghiệp vụ và phân quyền chính của hệ thống quản lý nhà hàng. "
        "Tài liệu được biên soạn theo schema hiện tại của dự án để hỗ trợ phát triển, kiểm thử, thuyết trình và bàn giao."
    )
    doc.add_paragraph(
        "Các bảng hạ tầng của Laravel như cache, jobs, password reset token, session kỹ thuật hoặc personal access token "
        "không được đưa vào phần mô tả chính để tài liệu tập trung hơn vào bài toán nghiệp vụ."
    )

    doc.add_heading("2. Quy ước thiết kế chung", level=1)
    for item in [
        "Các bảng nghiệp vụ mới đều có is_active để phục vụ ẩn dữ liệu hoặc ngừng sử dụng mềm mà không xóa cứng bản ghi.",
        "Các bảng nghiệp vụ mới đều có created_at và updated_at để theo dõi lịch sử tạo và cập nhật.",
        "Các trạng thái như trạng thái bàn, hóa đơn, thanh toán, dòng món... được chuẩn hóa bằng hằng số trong app/Common/Constants.",
        "Các bảng snapshot như cart_order_items, cart_order_item_options và invoice_items được thiết kế để giữ nguyên lịch sử tại thời điểm giao dịch.",
        "Bảng role_permissions là pivot cũ của module phân quyền nên không có is_active và timestamp; đây là ngoại lệ có chủ đích.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("3. Mô tả chi tiết từng bảng", level=1)

    migration_files = sorted(Path("database/migrations").glob("*.php"))
    table_columns = {}
    for path in migration_files:
        parsed = parse_migration(path)
        if parsed:
            table_columns[parsed[0]] = parsed[1]

    for index, table_name in enumerate(TARGET_TABLES, start=1):
        purpose, scenario = TABLE_META[table_name]
        columns = table_columns.get(table_name, [])

        doc.add_heading(f"3.{index}. Bảng {table_name}", level=2)
        doc.add_paragraph(f"Vì sao cần bảng này: {purpose}")
        doc.add_paragraph(f"Hoàn cảnh sử dụng điển hình: {scenario}")

        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        headers = ["Cột", "Kiểu / Ràng buộc", "Ý nghĩa", "Ví dụ / ghi chú"]
        for i, header in enumerate(headers):
            table.rows[0].cells[i].text = header

        for column in columns:
            row = table.add_row().cells
            for i, value in enumerate(column):
                row[i].text = value

        doc.add_paragraph("")

    doc.add_heading("4. Ghi chú triển khai", level=1)
    for item in [
        "Sơ đồ trực quan draw.io của cùng schema đã được tạo trong file docs/restaurant-er-diagram.drawio.xml.",
        "Nếu schema thay đổi về sau, chỉ cần cập nhật migration và chạy lại script này để sinh tài liệu Word mới.",
        "Tài liệu này ưu tiên tính dễ đọc cho người dùng và vẫn bám sát cấu trúc database hiện tại của dự án.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    return doc


def main() -> None:
    out_path = Path("docs/mo_ta_co_so_du_lieu_nha_hang.docx")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    build_document().save(out_path)
    print(out_path.resolve())


if __name__ == "__main__":
    main()
